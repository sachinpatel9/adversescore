"""Document generation — Phase 9 of the PSUR consolidation rebuild.

Renders a `ConsolidationResult` (Phase 6's pipeline output) into a
`python-docx` `Document` object structured as a trimmed subset of the ICH
E2C(R2) PBRER format: a cover/title page, Section 1 (Introduction),
Section 6 (Data in Summary Tabulations), Section 15 (Overview of Signals,
a real table), Section 16.1-16.3 (Signal and Risk Evaluation), a
consolidated placeholder note for the sections this tool does not populate,
a ranking-formula-version audit-trail statement, and a human-review
requirement statement.

Sections 1 and 6 are 100% deterministic — no LLM call. Sections 15's
introductory paragraph and 16.1-16.3's body text are produced by a single
narrow LLM call (`_narrate_signals`), grounded strictly in the structured
ranking/label data already computed by Phase 4/5 — never a full
tool-calling agent. This module deliberately does NOT import
`orchestrator.py` or `agent_tools.py`: it makes its own independent,
narrow completion call and reimplements the (pure, standalone) causal-
language flag rather than reaching into the agent-graph module.

Import-safety convention (matches `orchestrator.py`): importing this
module must never require `OPENAI_API_KEY`. Only *calling*
`generate_psur_document()` (via `_narrate_signals`) constructs the
`ChatOpenAI` model, and only on first call — the client is then cached at
module level for subsequent calls, mirroring `orchestrator.py`'s
`_cached_graph` pattern (tests substitute a fake model the same way, by
monkeypatching `document_generator._cached_narration_model` directly).

Returns the `python-docx` `Document` object itself — this module never
calls `.save()` and never returns bytes. Callers decide persistence.
"""
import re
from datetime import datetime, timezone
from typing import Optional

from docx import Document as _new_document
# `docx.Document` (above) is a *factory function*, not a class — using it as
# a type annotation would be misleading (and wrong under static type
# checkers). `docx.document.Document` (below) is the actual class returned
# by the factory; every type hint in this module uses that, matching the
# convention already established in this phase's own test files.
from docx.document import Document

from langchain_openai import ChatOpenAI

from .config import (
    DOCUMENT_NARRATION_TEMPERATURE,
    OPENAI_CHAT_MODEL,
    PBRER_DRAFT_MARKING_TEXT,
    PBRER_OMITTED_SECTIONS_NOTE,
    PBRER_PLACEHOLDER_TEXT,
    TOP_N_NARRATED_SIGNALS,
)
from .consolidation import ConsolidationResult
from .drug_identity import DrugIdentity
from .label_classifier import LabelClassificationResult
from .logger import get_logger, log_event
from .ranking import RankingResult

logger = get_logger("document_generator")


# ── Narration Guardrails (reimplemented standalone — see module docstring
# for why this does not import orchestrator.py's `_flag_causal_language`) ──

_CAUSAL_LANGUAGE_PATTERNS = (
    "causes", "cause of", "is known to cause", "caused by",
    "results in", "resulted in", "results from",
    "led to the", "leads to the", "due to the drug",
)


def _flag_causal_language_in_narration(text: str) -> bool:
    """Standalone equivalent of orchestrator.py's `_flag_causal_language`
    (Guardrail 1, code-enforced half) — log-only, never blocks or rewrites
    the narration. Reimplemented here (not imported) because this module
    must not depend on orchestrator.py."""
    lowered = text.lower()
    flagged = any(pattern in lowered for pattern in _CAUSAL_LANGUAGE_PATTERNS)
    if flagged:
        log_event(logger, "document_causal_language_flagged", snippet=text[:200])
    return flagged


_ALL_CAPS_TERM_RE = re.compile(r"\b[A-Z]{2,}[A-Z0-9\-]*\b")
_TITLE_CASE_MULTIWORD_RE = re.compile(r"\b[A-Z][a-z]+(?:\s+[A-Z][a-z]+)+\b")


def _check_narration_for_fabricated_entities(narrated_text: str, allowed_terms: set) -> list:
    """"No novel entities" fabrication check. `allowed_terms` must be the
    union of every ranked_signals[].symptom (uppercased) plus every
    drug_identity name variant (brand_names + generic_names +
    substance_names, uppercased) — see `_build_allowed_vocabulary()`.

    Scans `narrated_text` for capitalized multi-word phrases (e.g. "Cardiac
    Arrest") or all-caps terms (e.g. "IBUPROFEN") not present in
    `allowed_terms`. Returns a list of the flagged terms as they literally
    appeared in the text (empty list = clean). Logs a warning via
    `log_event` if any are found. Flag-only — never blocks/rejects the
    narration, matching Phase 8's causal-language flagging philosophy
    (log-only defense-in-depth, avoiding false positives from the LLM's
    natural narrative language, e.g. common acronyms like "PSUR" or "FAERS"
    will legitimately get flagged too since they are not in the allowed
    vocabulary — this is an accepted tradeoff of a flag-only check)."""
    if not narrated_text:
        return []

    flagged = []
    seen = set()
    candidates = _ALL_CAPS_TERM_RE.findall(narrated_text) + _TITLE_CASE_MULTIWORD_RE.findall(narrated_text)
    for term in candidates:
        normalized = term.strip().upper()
        if not normalized or normalized in allowed_terms or normalized in seen:
            continue
        seen.add(normalized)
        flagged.append(term)

    if flagged:
        log_event(logger, "narration_fabricated_entities_flagged", flagged_terms=flagged)
    return flagged


def _build_allowed_vocabulary(ranking: RankingResult, drug_identity: DrugIdentity) -> set:
    """Union of every ranked_signals[].symptom (already uppercased by
    ranking.py) plus every drug_identity name variant, uppercased —
    the ground-truth vocabulary the narration is allowed to reference."""
    allowed = {signal.symptom.upper() for signal in ranking.ranked_signals}
    for name in (
        list(drug_identity.brand_names)
        + list(drug_identity.generic_names)
        + list(drug_identity.substance_names)
        + [drug_identity.canonical_name]
    ):
        if name and name.strip():
            allowed.add(name.strip().upper())
    return allowed


# ── LLM Narration (single call, lazy client construction) ─────────────────
# `_cached_narration_model` mirrors orchestrator.py's `_cached_graph` —
# populated on first `_narrate_signals()` call, NOT at import time. Tests
# substitute a fake model by monkeypatching this module attribute directly.
_cached_narration_model = None

# The single LLM call is asked to produce two clearly delimited parts in one
# response (rather than two separate API calls) — a short introductory
# paragraph for Section 15's signal table, and a longer paragraph for
# Section 16.1-16.3's signal/risk evaluation framing. If the model does not
# honor the delimiter, `_split_narration` falls back to reusing the full
# text for both sections rather than dropping content.
_NARRATION_SPLIT_MARKER = "---SIGNAL_RISK_EVALUATION---"


def _get_narration_model():
    global _cached_narration_model
    if _cached_narration_model is None:
        _cached_narration_model = ChatOpenAI(
            model=OPENAI_CHAT_MODEL, temperature=DOCUMENT_NARRATION_TEMPERATURE)
    return _cached_narration_model


def _build_narration_prompt(
    ranking: RankingResult, drug_identity: DrugIdentity, label_summary: LabelClassificationResult
) -> str:
    """Builds the narration prompt from the ACTUAL ranked_signals data
    (symptom names, tiers, report counts, PRR point estimate) and
    label_summary counts — not a vague description. Itemizes at most
    TOP_N_NARRATED_SIGNALS signals by name for prompt-size discipline (the
    exported document's Section 15 TABLE itself is never capped — only this
    narration prompt's per-signal itemization is, mirroring the same
    TOP_N_NARRATED_SIGNALS convention Phase 8 uses for chat narration)."""
    top_signals = ranking.ranked_signals[:TOP_N_NARRATED_SIGNALS]
    signal_lines = []
    for signal in top_signals:
        prr_value = signal.prr_metrics.get("prr")
        ci_lower = signal.prr_metrics.get("ci_lower")
        signal_lines.append(
            f"- {signal.symptom}: seriousness={signal.seriousness_tier}, "
            f"strength_of_evidence={signal.strength_of_evidence_tier}, "
            f"reversibility={signal.reversibility_tier}, "
            f"public_health_impact={signal.public_health_tier}, "
            f"report_count={signal.report_count}, prr={prr_value}, ci_lower={ci_lower}"
        )
    signals_block = "\n".join(signal_lines) if signal_lines else "No signals were identified in this period."

    remaining = ranking.total_signals - len(top_signals)
    if remaining > 0:
        signals_block += f"\n(...and {remaining} additional lower-priority signal(s), all present in the full document table.)"

    name_variants = sorted({
        n.strip().upper() for n in (
            list(drug_identity.brand_names)
            + list(drug_identity.generic_names)
            + list(drug_identity.substance_names)
        ) if n and n.strip()
    })
    names_block = ", ".join(name_variants) if name_variants else drug_identity.canonical_name

    return f"""You are drafting narrative text for a PSUR (Periodic Safety Update
Report) aligned to ICH E2C(R2) PBRER format. You are given ONLY the
structured data below as ground truth. Follow these rules strictly:
- Narrate ONLY what is given below. Never introduce a drug, symptom, or
  event name that does not appear in the data below.
- Never assert or imply causation (never say "causes", "results in", "led
  to"). Use statistical-association language only (e.g. "associated
  with", "a signal was detected for").
- Always frame the ranking/PRR data as a statistical signal requiring
  clinical validation — NOT a completed clinical evaluation and NOT proof
  of causation. Disproportionality statistics (like PRR) are evidentiary
  input only.

DRUG: {drug_identity.canonical_name} (name variants: {names_block})

LABEL STATUS SUMMARY (baseline safety concerns): {label_summary.total_symptoms} total
symptoms classified — {label_summary.labeled_count} LABELED (already on the approved
label), {label_summary.unlabeled_count} UNLABELED, {label_summary.unknown_count} UNKNOWN label status.

RANKED SIGNALS ({ranking.total_signals} total, top {len(top_signals)} itemized below):
{signals_block}

Produce your response in exactly two parts, separated by a line containing
ONLY the literal text {_NARRATION_SPLIT_MARKER}

PART 1 (2-4 sentences): A short introductory paragraph contextualizing the
signal table for a PSUR reader — what the table shows and how it is
organized (tiered, not a single composite score).

PART 2 (1-3 short paragraphs): Text for the "Signal and Risk Evaluation"
section covering (a) the baseline safety concerns implied by the label
status counts above, and (b) framing of the ranking/PRR data as
evidentiary input requiring clinical validation, not a finished
evaluation. Explicitly state this is not a completed clinical evaluation.
"""


def _narrate_signals(
    ranking: RankingResult, drug_identity: DrugIdentity, label_summary: LabelClassificationResult
) -> str:
    """Single LLM call. Prompt provides ONLY the structured ranking/label
    data as ground truth. Instructs: narrate only what's given, never
    introduce a drug/event not in the provided data, never assert
    causation, always frame as statistical signal requiring clinical
    validation (not causation, not a finished evaluation)."""
    prompt = _build_narration_prompt(ranking, drug_identity, label_summary)
    model = _get_narration_model()
    response = model.invoke(prompt)
    content = getattr(response, "content", None)
    text = content if isinstance(content, str) else (str(content) if content else "")
    log_event(logger, "signal_narration_generated", total_signals=ranking.total_signals)
    return text


def _split_narration(narrated_text: str) -> tuple:
    """Splits the single LLM narration response into (intro_text,
    risk_evaluation_text) on `_NARRATION_SPLIT_MARKER`. Defensive fallback
    if the model didn't honor the delimiter format: reuse the full text for
    both sections rather than dropping content or crashing."""
    if _NARRATION_SPLIT_MARKER in narrated_text:
        intro, risk = narrated_text.split(_NARRATION_SPLIT_MARKER, 1)
        return intro.strip(), risk.strip()
    stripped = narrated_text.strip()
    return stripped, stripped


def _get_signal_narration(result: ConsolidationResult) -> tuple:
    """Runs the single LLM narration call, applies the flag-only guardrail
    checks (fabrication + causal language — both log-only, never block),
    and returns (intro_text, risk_evaluation_text)."""
    narrated_text = _narrate_signals(result.ranking, result.drug_identity, result.ranking.label_summary)

    allowed_terms = _build_allowed_vocabulary(result.ranking, result.drug_identity)
    _check_narration_for_fabricated_entities(narrated_text, allowed_terms)
    _flag_causal_language_in_narration(narrated_text)

    return _split_narration(narrated_text)


# ── Deterministic Guardrail Text Guarantees (code-side, idempotent) ───────

# Two independently-required framing components (mirrors the fix already
# applied to orchestrator.py's `_ensure_completeness_statement`, which found
# that checking a single generic marker in isolation — e.g. "retrieved" —
# let a response that never actually disclosed the real content silently
# suppress the deterministic guarantee). "clinical validation" alone is too
# generic a phrase to prove the narration stated BOTH that the ranking is
# evidentiary input only AND that it is not a completed clinical evaluation
# — a response like "further clinical validation would help" could satisfy
# a single-marker check without ever conveying either required framing, so
# both groups must be independently satisfied before the disclaimer is
# skipped.
_EVIDENTIARY_FRAMING_MARKERS = ("evidentiary input",)
_NOT_COMPLETED_FRAMING_MARKERS = ("not a completed", "not a finished", "not a validated safety signal")
_EVIDENTIARY_INPUT_DISCLAIMER = (
    "This ranking reflects a statistical disproportionality signal derived from FAERS "
    "spontaneous reports (evidentiary input only) and is not a validated safety signal or "
    "a completed clinical evaluation; per ICH E2C(R2), a disproportionality statistic "
    "requires further clinical validation before it constitutes an evaluated risk."
)


def _ensure_evidentiary_input_disclaimer(text: str) -> str:
    """Deterministic, idempotent code-side guarantee that Section 16's text
    explicitly states the evidentiary-input/not-a-finished-evaluation
    framing, mirroring this repo's established guardrail-postprocessing
    pattern (orchestrator.py's `_ensure_*` helpers) — the LLM is already
    instructed to say this via the prompt, but the document must never
    depend solely on the model actually following that instruction.

    Requires BOTH the evidentiary-input framing AND the not-a-completed-
    evaluation framing to already be present before skipping the append —
    see module-level comment above `_EVIDENTIARY_FRAMING_MARKERS` for why a
    single generic marker match is not sufficient evidence either
    disclosure was actually made."""
    lowered = text.lower()
    has_evidentiary_framing = any(marker in lowered for marker in _EVIDENTIARY_FRAMING_MARKERS)
    has_not_completed_framing = any(marker in lowered for marker in _NOT_COMPLETED_FRAMING_MARKERS)
    if has_evidentiary_framing and has_not_completed_framing:
        return text
    if not text:
        return _EVIDENTIARY_INPUT_DISCLAIMER
    return text + "\n\n" + _EVIDENTIARY_INPUT_DISCLAIMER


_HUMAN_REVIEW_STATEMENT = (
    "Human Review Requirement: This draft document and the underlying signal ranking "
    "require review and sign-off by a qualified pharmacovigilance and/or clinical "
    "professional before any regulatory submission or clinical decision-making use. "
    "AdverseScore does not perform this review and does not substitute for expert "
    "clinical judgment."
)


# ── Document Section Builders (deterministic, no LLM) ─────────────────────

def _apply_header_footer(document: Document) -> None:
    """Sets the running header AND footer to PBRER_DRAFT_MARKING_TEXT on
    every section (here, the document's single default section) — a real
    docx header/footer paragraph, structurally distinct from body text."""
    for section in document.sections:
        header_paragraph = section.header.paragraphs[0] if section.header.paragraphs else section.header.add_paragraph()
        header_paragraph.text = PBRER_DRAFT_MARKING_TEXT

        footer_paragraph = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
        footer_paragraph.text = PBRER_DRAFT_MARKING_TEXT


def _add_cover_page(document: Document, result: ConsolidationResult) -> None:
    generated_at = datetime.now(timezone.utc)
    document.add_heading(result.drug_identity.canonical_name, level=0)
    document.add_paragraph("Periodic Safety Update Report (PSUR) — Signal Consolidation Draft")
    document.add_paragraph(f"Reporting period: {result.period} "
                            f"({result.retrieval.period_start} to {result.retrieval.period_end})")
    document.add_paragraph(f"Generated: {generated_at.strftime('%Y-%m-%d %H:%M:%S UTC')}")
    document.add_paragraph(PBRER_DRAFT_MARKING_TEXT)
    document.add_page_break()


def _add_section_1_introduction(document: Document, result: ConsolidationResult) -> None:
    identity = result.drug_identity
    document.add_heading("Section 1 — Introduction", level=1)

    document.add_paragraph(f"Canonical name: {identity.canonical_name}")
    document.add_paragraph(f"Brand name(s): {', '.join(identity.brand_names) or 'none identified'}")
    document.add_paragraph(f"Generic name(s): {', '.join(identity.generic_names) or 'none identified'}")

    if identity.market_authorization_date is not None:
        document.add_paragraph(
            f"Market authorization date: {identity.market_authorization_date.isoformat()} "
            f"(source: {identity.market_authorization_date_source})")
    else:
        document.add_paragraph(
            "Market authorization date: not available. This is expected for OTC monograph "
            "drugs, which are approved via a monograph pathway rather than an NDA/ANDA/BLA "
            f"and therefore have no drugsfda.json entry (source note: "
            f"{identity.market_authorization_date_source}).")

    if result.pharm_class:
        document.add_paragraph(f"Pharmacologic class: {result.pharm_class}")
    else:
        document.add_paragraph("Pharmacologic class: not available (class discovery unsuccessful).")

    document.add_paragraph(f"Reporting period: {result.period}")


def _add_section_6_data_summary(document: Document, result: ConsolidationResult) -> None:
    retrieval = result.retrieval
    dedup = result.dedup
    document.add_heading("Section 6 — Data in Summary Tabulations", level=1)

    document.add_paragraph(
        f"Retrieved reports: {retrieval.total_retrieved_count} of an estimated "
        f"{retrieval.total_estimated_count} matching FAERS reports.")

    if retrieval.any_chunk_truncated:
        truncated_labels = [c.label for c in retrieval.chunks if c.truncated]
        document.add_paragraph(
            "Retrieval completeness: one or more retrieval chunks were truncated by "
            f"openFDA pagination limits. Truncated chunk(s): {', '.join(truncated_labels) or 'unspecified'}.")
    else:
        document.add_paragraph("Retrieval completeness: no chunk truncation occurred.")

    document.add_paragraph(
        "Deduplication methodology: of "
        f"{dedup.total_input_count} input reports, {dedup.total_output_count} remained after "
        f"deduplication ({dedup.removed_by_exact_id} removed by exact safetyreportid match, "
        f"{dedup.removed_by_heuristic} removed by heuristic match on drug names + symptoms + date).")


_SIGNAL_TABLE_HEADERS = (
    "Symptom (MedDRA PT)",
    "Seriousness Tier",
    "Strength of Evidence Tier",
    "Reversibility Tier",
    "Public Health Tier",
    "Report Count",
    "PRR (95% CI lower bound)",
)


def _add_signal_table(document: Document, result: ConsolidationResult) -> None:
    ranked_signals = result.ranking.ranked_signals
    if not ranked_signals:
        document.add_paragraph("No signals were identified in this reporting period.")
        return

    table = document.add_table(rows=1, cols=len(_SIGNAL_TABLE_HEADERS))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for i, header_text in enumerate(_SIGNAL_TABLE_HEADERS):
        header_cells[i].text = header_text

    for signal in ranked_signals:
        prr_value = signal.prr_metrics.get("prr")
        ci_lower = signal.prr_metrics.get("ci_lower")
        row_cells = table.add_row().cells
        row_cells[0].text = signal.symptom
        row_cells[1].text = signal.seriousness_tier
        row_cells[2].text = signal.strength_of_evidence_tier
        row_cells[3].text = signal.reversibility_tier
        row_cells[4].text = signal.public_health_tier
        row_cells[5].text = str(signal.report_count)
        row_cells[6].text = f"{prr_value} ({ci_lower})" if prr_value is not None else "N/A"


def _add_section_15_overview_of_signals(document: Document, result: ConsolidationResult, intro_text: str) -> None:
    document.add_heading("Section 15 — Overview of Signals", level=1)
    if intro_text:
        document.add_paragraph(intro_text)
    document.add_paragraph(
        f"Full ranked signal list ({result.ranking.total_signals} signal(s)), formula version "
        f"{result.ranking.formula_version}:")
    _add_signal_table(document, result)


def _add_section_16_signal_risk_evaluation(document: Document, result: ConsolidationResult, risk_text: str) -> None:
    document.add_heading("Section 16.1–16.3 — Signal and Risk Evaluation", level=1)
    document.add_heading("16.1 Summary of Safety Concerns / 16.2–16.3 Signal Evaluation", level=2)

    label_summary = result.ranking.label_summary
    document.add_paragraph(
        f"Baseline label status across {label_summary.total_symptoms} classified symptom(s): "
        f"{label_summary.labeled_count} LABELED, {label_summary.unlabeled_count} UNLABELED, "
        f"{label_summary.unknown_count} UNKNOWN.")

    body_text = _ensure_evidentiary_input_disclaimer(risk_text or "")
    document.add_paragraph(body_text)


def _add_omitted_sections_note(document: Document) -> None:
    document.add_heading("Sections Not Populated in This Draft", level=1)
    document.add_paragraph(f"{PBRER_OMITTED_SECTIONS_NOTE} {PBRER_PLACEHOLDER_TEXT}")


def _add_formula_version_statement(document: Document, result: ConsolidationResult) -> None:
    document.add_paragraph(f"Ranking formula version: {result.formula_version}")


def _add_human_review_statement(document: Document) -> None:
    document.add_paragraph(_HUMAN_REVIEW_STATEMENT)


# ── Entry Point ────────────────────────────────────────────────────────────

def generate_psur_document(result: ConsolidationResult) -> Document:
    """Renders `result` (a Phase 6 `ConsolidationResult`) into a python-docx
    `Document` object structured as a trimmed ICH E2C(R2) PBRER subset (see
    module docstring for the section list). Returns the `Document` object
    itself — never saves to a path or returns bytes; callers decide
    `.save(path)` or a `BytesIO` write themselves.

    Sections 1 and 6 are fully deterministic. Section 15's intro paragraph
    and Section 16.1-16.3's body text come from a single LLM narration call
    (`_narrate_signals`), guardrailed by flag-only fabrication and
    causal-language checks (log, never block) plus a deterministic
    idempotent disclaimer guarantee for the evidentiary-input framing.
    """
    document = _new_document()

    _apply_header_footer(document)
    _add_cover_page(document, result)
    _add_section_1_introduction(document, result)
    _add_section_6_data_summary(document, result)

    intro_text, risk_text = _get_signal_narration(result)
    _add_section_15_overview_of_signals(document, result, intro_text)
    _add_section_16_signal_risk_evaluation(document, result, risk_text)

    _add_omitted_sections_note(document)
    _add_formula_version_statement(document, result)
    _add_human_review_statement(document)

    log_event(
        logger, "psur_document_generated",
        drug=result.drug_identity.canonical_name, period=result.period,
        total_signals=result.ranking.total_signals,
    )

    return document
