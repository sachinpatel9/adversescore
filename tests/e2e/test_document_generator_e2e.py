"""
End-to-end integration test for document_generator.py (Phase 9) — runs the
full consolidate_psur() pipeline against the live openFDA API, then generates
a real PSUR .docx document via a live OpenAI narration call.
"""
import pytest

pytestmark = pytest.mark.e2e
from conftest import SKIP_NO_FDA, SKIP_NO_OPENAI
from docx.document import Document

from adverse_score.config import PBRER_DRAFT_MARKING_TEXT, PBRER_PLACEHOLDER_TEXT
from adverse_score.consolidation import consolidate_psur, ConsolidationResult
from adverse_score.document_generator import generate_psur_document


@SKIP_NO_FDA
@SKIP_NO_OPENAI
class TestDocumentGeneratorE2E:
    def test_generate_psur_document_keytruda_6mo(self):
        """Live consolidate_psur("KEYTRUDA", "6mo") -> generate_psur_document(result)
        against real FAERS data and a real LLM narration call. Narrated prose
        content varies live — only structural invariants are asserted."""
        result = consolidate_psur("KEYTRUDA", "6mo")
        assert isinstance(result, ConsolidationResult)

        document = generate_psur_document(result)
        assert isinstance(document, Document)

        headings = [
            p.text for p in document.paragraphs
            if p.style.name.startswith("Heading") or p.style.name == "Title"
        ]
        joined = "\n".join(headings)
        assert "Introduction" in joined
        assert "Data in Summary Tabulations" in joined
        assert "Overview of Signals" in joined
        assert "Signal and Risk Evaluation" in joined

        paragraph_texts = [p.text for p in document.paragraphs]
        assert any(PBRER_PLACEHOLDER_TEXT in t for t in paragraph_texts)
        assert PBRER_DRAFT_MARKING_TEXT in paragraph_texts

        section = document.sections[0]
        assert PBRER_DRAFT_MARKING_TEXT in [p.text for p in section.header.paragraphs]
        assert PBRER_DRAFT_MARKING_TEXT in [p.text for p in section.footer.paragraphs]

        assert any(result.formula_version in t for t in paragraph_texts)

        # KEYTRUDA over a 6-month PSUR window reliably has ranked signals — this
        # proves real data flowed all the way through to the exported table.
        assert result.ranking.total_signals > 0
        assert len(document.tables) == 1
        table = document.tables[0]
        assert len(table.rows) == len(result.ranking.ranked_signals) + 1  # + header row
