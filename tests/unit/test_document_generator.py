"""
Unit tests for document_generator.py (Phase 9) — deterministic PBRER-subset
.docx generation plus the single-call LLM narration and its flag-only
guardrail checks.

The LLM is never called live: `langchain_core.language_models.fake_chat_models.
FakeListChatModel` (a plain, non-tool-calling fake — this phase's LLM call is
a single completion, not an agent) is installed via monkeypatching
`document_generator._cached_narration_model` directly, mirroring
`test_orchestrator.py`'s `fake_graph` fixture pattern for `_cached_graph`.
"""
from datetime import date

import pytest
from docx.document import Document
from langchain_core.language_models.fake_chat_models import FakeListChatModel

from adverse_score import document_generator
from adverse_score.config import (
    PBRER_DRAFT_MARKING_TEXT,
    PBRER_OMITTED_SECTIONS_NOTE,
    PBRER_PLACEHOLDER_TEXT,
)
from adverse_score.consolidation import ConsolidationResult
from adverse_score.deduplication import DedupResult
from adverse_score.document_generator import (
    _EVIDENTIARY_INPUT_DISCLAIMER,
    _NARRATION_SPLIT_MARKER,
    _check_narration_for_fabricated_entities,
    _ensure_evidentiary_input_disclaimer,
    _split_narration,
    generate_psur_document,
)
from adverse_score.drug_identity import DrugIdentity
from adverse_score.fda_client import ChunkResult, PSURRetrievalResult
from adverse_score.label_classifier import LabelClassificationResult
from adverse_score.ranking import RankedSignal, RankingResult


# ── Fixtures / builders ─────────────────────────────────────────────────────

def _drug_identity(**overrides):
    defaults = dict(
        canonical_name="KEYTRUDA",
        brand_names=["KEYTRUDA"],
        generic_names=["PEMBROLIZUMAB"],
        substance_names=["PEMBROLIZUMAB"],
        application_numbers=["BLA125514"],
        market_authorization_date=date(2014, 9, 4),
        market_authorization_date_source="drugsfda.json:BLA125514",
        resolution_confidence="EXACT",
    )
    defaults.update(overrides)
    return DrugIdentity(**defaults)


def _ranked_signal(i, **overrides):
    defaults = dict(
        symptom=f"SIGNAL_{i}", rank=i + 1, seriousness_tier="NON_SERIOUS",
        strength_of_evidence_tier="WEAK_UNLABELED", reversibility_tier="UNKNOWN",
        public_health_tier="LOW",
        prr_metrics={"prr": 1.5, "ci_lower": 0.5, "signal_detected": False,
                     "target_symptom": f"SIGNAL_{i}", "drug_cases": 5,
                     "class_cases": 10, "label_status": "UNLABELED"},
        report_count=5,
    )
    defaults.update(overrides)
    return RankedSignal(**defaults)


def _consolidation_result(n_signals=3, any_chunk_truncated=False, truncated_chunk_labels=None,
                          identity=None, retrieved_count=100, estimated_count=120,
                          total_input_count=100, total_output_count=90,
                          removed_by_exact_id=8, removed_by_heuristic=2,
                          pharm_class="PD-1 inhibitors", **overrides):
    identity = identity or _drug_identity()
    truncated_chunk_labels = truncated_chunk_labels or (["chunk-0-20240101-to-20240401"] if any_chunk_truncated else [])

    chunks = [
        ChunkResult(
            label="chunk-0-20240101-to-20240401", start_date="20240101", end_date="20240401",
            reports=[], retrieved_count=retrieved_count, estimated_total_count=estimated_count,
            truncated="chunk-0-20240101-to-20240401" in truncated_chunk_labels, error=None,
        )
    ]
    retrieval = PSURRetrievalResult(
        canonical_query_variants=[identity.canonical_name], period_start="20240101", period_end="20240701",
        chunks=chunks, reports=[], total_retrieved_count=retrieved_count, total_estimated_count=estimated_count,
        any_chunk_truncated=any_chunk_truncated, used_fallback_anchor=False, fallback_reason=None,
    )
    dedup = DedupResult(
        reports=[], total_input_count=total_input_count, total_output_count=total_output_count,
        removed_by_exact_id=removed_by_exact_id, removed_by_heuristic=removed_by_heuristic, audit_trail=[],
    )
    ranked_signals = [_ranked_signal(i) for i in range(n_signals)]
    label_summary = LabelClassificationResult(
        statuses={s.symptom: "UNLABELED" for s in ranked_signals},
        total_symptoms=n_signals, labeled_count=0, unlabeled_count=n_signals, unknown_count=0,
    )
    ranking = RankingResult(
        ranked_signals=ranked_signals, label_summary=label_summary,
        formula_version="1.0", total_signals=n_signals,
    )
    defaults = dict(
        drug_identity=identity, period="6mo", retrieval=retrieval, dedup=dedup,
        ranking=ranking, pharm_class=pharm_class, class_counts_available=True,
        formula_version="1.0",
    )
    defaults.update(overrides)
    return ConsolidationResult(**defaults)


_CANNED_NARRATION = (
    "This table summarizes the ranked adverse event signals identified for the "
    "drug across four evidentiary tiers, ordered by descending priority."
    f"\n{_NARRATION_SPLIT_MARKER}\n"
    "Baseline label status indicates most signals are currently unlabeled. "
    "This ranking reflects a statistical disproportionality signal only and is "
    "not a completed clinical evaluation; further clinical validation is required."
)


@pytest.fixture
def fake_narration_model(monkeypatch):
    """Installs a FakeListChatModel as document_generator's cached narration
    model, bypassing ChatOpenAI construction entirely. Returns a setter the
    test calls with its scripted response text (defaults to a clean, on-format
    canned narration if never called)."""
    def _install(response_text=_CANNED_NARRATION):
        model = FakeListChatModel(responses=[response_text])
        monkeypatch.setattr(document_generator, "_cached_narration_model", model)
        return model
    return _install


def _all_paragraph_texts(document: Document) -> list:
    return [p.text for p in document.paragraphs]


def _all_text(document: Document) -> str:
    parts = list(_all_paragraph_texts(document))
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


# ── Structural: headers/body text present ───────────────────────────────────

class TestDocumentStructure:
    def test_generate_psur_document_returns_document(self, fake_narration_model):
        fake_narration_model()
        result = _consolidation_result()
        document = generate_psur_document(result)
        assert isinstance(document, Document)

    def test_required_section_headers_present(self, fake_narration_model):
        fake_narration_model()
        result = _consolidation_result()
        document = generate_psur_document(result)
        headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"]
        joined = "\n".join(headings)
        assert "Introduction" in joined
        assert "Data in Summary Tabulations" in joined
        assert "Overview of Signals" in joined
        assert "Signal and Risk Evaluation" in joined

    def test_consolidated_placeholder_note_present(self, fake_narration_model):
        fake_narration_model()
        result = _consolidation_result()
        document = generate_psur_document(result)
        texts = _all_paragraph_texts(document)
        expected = f"{PBRER_OMITTED_SECTIONS_NOTE} {PBRER_PLACEHOLDER_TEXT}"
        assert expected in texts

    def test_omitted_sections_note_covers_appendices(self, fake_narration_model):
        """Regression: the built sections are 1, 6, 15, 16.1-16.3 — every
        other ICH E2C(R2) PBRER section (including Section 20, Appendices)
        must be accounted for by the omitted-sections note. Section 20 was
        previously missing from both the built set and the note's stated
        ranges, silently falling through both buckets."""
        assert "20" in PBRER_OMITTED_SECTIONS_NOTE

    def test_placeholder_text_appears_verbatim(self, fake_narration_model):
        """Byte-exact acceptance criterion: PBRER_PLACEHOLDER_TEXT must appear
        character-for-character in some paragraph, not paraphrased."""
        fake_narration_model()
        result = _consolidation_result()
        document = generate_psur_document(result)
        assert any(PBRER_PLACEHOLDER_TEXT in p.text for p in document.paragraphs)
        # Exact match, not just substring within a longer sentence somewhere else
        exact_matches = [p.text for p in document.paragraphs if p.text == f"{PBRER_OMITTED_SECTIONS_NOTE} {PBRER_PLACEHOLDER_TEXT}"]
        assert len(exact_matches) == 1

    def test_draft_marking_in_body_paragraph(self, fake_narration_model):
        fake_narration_model()
        result = _consolidation_result()
        document = generate_psur_document(result)
        assert any(p.text == PBRER_DRAFT_MARKING_TEXT for p in document.paragraphs)

    def test_draft_marking_in_header_and_footer(self, fake_narration_model):
        """Structurally distinct check: header/footer paragraphs, not body text."""
        fake_narration_model()
        result = _consolidation_result()
        document = generate_psur_document(result)
        section = document.sections[0]
        header_texts = [p.text for p in section.header.paragraphs]
        footer_texts = [p.text for p in section.footer.paragraphs]
        assert PBRER_DRAFT_MARKING_TEXT in header_texts
        assert PBRER_DRAFT_MARKING_TEXT in footer_texts

    def test_formula_version_present(self, fake_narration_model):
        fake_narration_model()
        result = _consolidation_result(formula_version="7.3")
        document = generate_psur_document(result)
        assert any("7.3" in p.text for p in document.paragraphs)

    def test_human_review_statement_present(self, fake_narration_model):
        fake_narration_model()
        result = _consolidation_result()
        document = generate_psur_document(result)
        assert any("human review" in p.text.lower() or "review and sign-off" in p.text.lower()
                   for p in document.paragraphs)


# ── Completeness data matches input ─────────────────────────────────────────

class TestCompletenessData:
    def test_retrieved_and_estimated_counts_present(self, fake_narration_model):
        fake_narration_model()
        result = _consolidation_result(retrieved_count=4321, estimated_count=9876)
        document = generate_psur_document(result)
        text = _all_text(document)
        assert "4321" in text
        assert "9876" in text

    def test_truncation_flag_and_chunk_label_present_when_truncated(self, fake_narration_model):
        fake_narration_model()
        result = _consolidation_result(any_chunk_truncated=True)
        document = generate_psur_document(result)
        text = _all_text(document)
        assert "truncat" in text.lower()
        assert "chunk-0-20240101-to-20240401" in text

    def test_no_truncation_statement_when_not_truncated(self, fake_narration_model):
        fake_narration_model()
        result = _consolidation_result(any_chunk_truncated=False)
        document = generate_psur_document(result)
        text = _all_text(document)
        assert "no chunk truncation occurred" in text.lower()

    def test_dedup_methodology_numbers_present(self, fake_narration_model):
        fake_narration_model()
        result = _consolidation_result(
            total_input_count=555, total_output_count=444,
            removed_by_exact_id=77, removed_by_heuristic=34,
        )
        document = generate_psur_document(result)
        text = _all_text(document)
        assert "555" in text
        assert "444" in text
        assert "77" in text
        assert "34" in text


# ── Section 15 signal table ─────────────────────────────────────────────────

class TestSignalTable:
    def test_table_is_real_docx_table_with_correct_row_count(self, fake_narration_model):
        fake_narration_model()
        result = _consolidation_result(n_signals=3)
        document = generate_psur_document(result)
        assert len(document.tables) == 1
        table = document.tables[0]
        # 1 header row + 3 data rows
        assert len(table.rows) == 4
        assert len(table.columns) == 7

    def test_table_row_count_matches_two_signals(self, fake_narration_model):
        fake_narration_model()
        result = _consolidation_result(n_signals=2)
        document = generate_psur_document(result)
        table = document.tables[0]
        assert len(table.rows) == 3  # header + 2

    def test_table_contains_all_ranked_signals_not_capped(self, fake_narration_model):
        """The scope doc explicitly requires the FULL ranked list in the
        exported document, unlike the top-N conversational narration cap."""
        fake_narration_model()
        n_signals = 25  # deliberately exceeds TOP_N_NARRATED_SIGNALS=20
        result = _consolidation_result(n_signals=n_signals)
        document = generate_psur_document(result)
        table = document.tables[0]
        assert len(table.rows) == n_signals + 1

    def test_table_cell_values_match_signal_fields(self, fake_narration_model):
        fake_narration_model()
        result = _consolidation_result(n_signals=1)
        document = generate_psur_document(result)
        table = document.tables[0]
        data_row = table.rows[1].cells
        signal = result.ranking.ranked_signals[0]
        assert data_row[0].text == signal.symptom
        assert data_row[1].text == signal.seriousness_tier
        assert data_row[2].text == signal.strength_of_evidence_tier
        assert data_row[3].text == signal.reversibility_tier
        assert data_row[4].text == signal.public_health_tier
        assert data_row[5].text == str(signal.report_count)

    def test_empty_ranked_signals_handled_gracefully(self, fake_narration_model):
        fake_narration_model()
        result = _consolidation_result(n_signals=0)
        document = generate_psur_document(result)
        # No table should be created for an empty signal list.
        assert len(document.tables) == 0
        text = _all_text(document)
        assert "no signals" in text.lower()


# ── OTC / missing market authorization date ────────────────────────────────

class TestOTCMarketAuthorizationDate:
    def test_none_market_authorization_date_handled_gracefully(self, fake_narration_model):
        fake_narration_model()
        identity = _drug_identity(
            market_authorization_date=None,
            market_authorization_date_source="no AP-status submission found in drugsfda.json",
            resolution_confidence="PARTIAL",
        )
        result = _consolidation_result(identity=identity)
        # Must not raise.
        document = generate_psur_document(result)
        text = _all_text(document)
        assert "not available" in text.lower()


# ── Fabrication guardrail ────────────────────────────────────────────────

class TestFabricationCheck:
    def test_flags_terms_not_in_allowed_vocabulary(self):
        allowed = {"NAUSEA", "KEYTRUDA"}
        text = "The patient also experienced Severe Hepatitis after taking Ibuprofen."
        flagged = _check_narration_for_fabricated_entities(text, allowed)
        assert len(flagged) > 0

    def test_clean_text_referencing_only_allowed_terms_returns_empty(self):
        allowed = {"NAUSEA", "KEYTRUDA", "FATIGUE"}
        text = "A signal was detected for nausea and fatigue associated with keytruda in this period."
        flagged = _check_narration_for_fabricated_entities(text, allowed)
        assert flagged == []

    def test_all_caps_term_not_allowed_is_flagged(self):
        allowed = {"NAUSEA"}
        text = "The reports also mention IBUPROFEN as a concomitant medication."
        flagged = _check_narration_for_fabricated_entities(text, allowed)
        assert any("IBUPROFEN" in f.upper() for f in flagged)

    def test_allowed_all_caps_term_not_flagged(self):
        allowed = {"NAUSEA", "IBUPROFEN"}
        text = "IBUPROFEN reports mention NAUSEA frequently."
        flagged = _check_narration_for_fabricated_entities(text, allowed)
        assert flagged == []

    def test_empty_text_returns_empty_list(self):
        assert _check_narration_for_fabricated_entities("", {"NAUSEA"}) == []


# ── Narration splitting ─────────────────────────────────────────────────

class TestSplitNarration:
    def test_splits_on_marker(self):
        text = f"intro paragraph{_NARRATION_SPLIT_MARKER}risk evaluation paragraph"
        intro, risk = _split_narration(text)
        assert intro == "intro paragraph"
        assert risk == "risk evaluation paragraph"

    def test_falls_back_to_full_text_when_marker_missing(self):
        text = "a single undivided narration block"
        intro, risk = _split_narration(text)
        assert intro == text
        assert risk == text


# ── Narration content flows into the document ──────────────────────────────

class TestNarrationIntegration:
    def test_narrated_intro_text_appears_in_document(self, fake_narration_model):
        fake_narration_model("Intro sentence about the table."
                              f"{_NARRATION_SPLIT_MARKER}"
                              "Risk evaluation sentence, evidentiary input, not a completed evaluation.")
        result = _consolidation_result()
        document = generate_psur_document(result)
        text = _all_text(document)
        assert "Intro sentence about the table." in text
        assert "Risk evaluation sentence" in text

    def test_evidentiary_disclaimer_appended_when_llm_omits_it(self, fake_narration_model):
        """Deterministic code-side guarantee: even if the LLM's Part 2 text
        doesn't mention evidentiary-input framing, the document must still
        state it (idempotent guardrail, matching orchestrator.py's pattern)."""
        fake_narration_model(f"Intro.{_NARRATION_SPLIT_MARKER}Some risk text with no disclaimer language.")
        result = _consolidation_result()
        document = generate_psur_document(result)
        text = _all_text(document).lower()
        assert "evidentiary input" in text or "clinical validation" in text


# ── _ensure_evidentiary_input_disclaimer — direct unit coverage ────────────
# Mirrors orchestrator.py's `_ensure_completeness_statement` fix: a single
# generic marker match (e.g. "clinical validation") must never be enough on
# its own to silently suppress the real disclaimer, since a narration could
# use that phrase without ever conveying either required framing component.

class TestEnsureEvidentiaryInputDisclaimer:
    def test_no_markers_present_appends_disclaimer(self):
        text = "The signals identified are ranked by tier."
        result = _ensure_evidentiary_input_disclaimer(text)
        assert result == text + "\n\n" + _EVIDENTIARY_INPUT_DISCLAIMER

    def test_generic_clinical_validation_marker_alone_does_not_suppress_disclaimer(self):
        """Regression: 'clinical validation' alone (without the evidentiary-
        input framing or the not-a-completed-evaluation framing) previously
        satisfied a single-marker OR check and silently suppressed the real
        disclosure. It must not be sufficient on its own."""
        text = "We recommend further clinical validation to confirm the finding."
        result = _ensure_evidentiary_input_disclaimer(text)
        assert result == text + "\n\n" + _EVIDENTIARY_INPUT_DISCLAIMER

    def test_only_evidentiary_framing_without_not_completed_framing_still_appends(self):
        text = "This is evidentiary input for the reader's consideration."
        result = _ensure_evidentiary_input_disclaimer(text)
        assert result == text + "\n\n" + _EVIDENTIARY_INPUT_DISCLAIMER

    def test_only_not_completed_framing_without_evidentiary_framing_still_appends(self):
        text = "This is not a completed clinical evaluation of the drug."
        result = _ensure_evidentiary_input_disclaimer(text)
        assert result == text + "\n\n" + _EVIDENTIARY_INPUT_DISCLAIMER

    def test_both_framings_present_does_not_duplicate_disclaimer(self):
        text = ("This reflects evidentiary input only and is not a completed "
                "clinical evaluation.")
        result = _ensure_evidentiary_input_disclaimer(text)
        assert result == text

    def test_empty_text_returns_bare_disclaimer(self):
        assert _ensure_evidentiary_input_disclaimer("") == _EVIDENTIARY_INPUT_DISCLAIMER
